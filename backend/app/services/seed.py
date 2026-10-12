from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import numpy as np
import pandas as pd
from faker import Faker
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import func, select, text

from ..db import ROOT, Base
from ..models import (
    Cobranca,
    Configuracao,
    Documento,
    Evento,
    Fornecedor,
    Justificativa,
    Pendencia,
    Projeto,
    Usuario,
)
from .rules import DEFAULT_CONFIG, audit, evaluate, refresh_document

AREAS = ["Operações", "Tecnologia", "Engenharia", "Suprimentos", "Administrativo"]
REFERENCE = date(2026, 10, 11)


def seed(db, export=True):
    if db.scalar(select(Usuario.id).limit(1)):
        return
    rng = np.random.default_rng(42)
    fake = Faker("pt_BR")
    fake.seed_instance(42)
    db.add(Configuracao(id=1, parametros=DEFAULT_CONFIG))
    for i in range(1, 66):
        db.add(
            Usuario(
                id=i,
                nome=f"Pessoa fictícia {i:02d} · {fake.first_name()}",
                email=f"usuario{i}@atlas.example",
                area=AREAS[(i - 1) % 5],
                perfil="administrador" if i == 1 else "revisor" if i <= 6 else "operador",
            )
        )
    db.flush()
    names = [
        "Horizonte",
        "Aurora",
        "Nexus",
        "Vértice",
        "Órbita",
        "Ponte",
        "Atlas",
        "Prisma",
        "Conecta",
    ]
    for i in range(1, 46):
        db.add(
            Projeto(
                id=i,
                codigo=f"PRJ-{i:03d}",
                nome=f"{names[(i - 1) % 9]} / Unidade {(i - 1) // 9 + 1}",
                descricao="Projeto demonstrativo de Atlas Operações. Todos os dados são fictícios.",
                centro_custo=f"CC-{1000 + i}",
                area=AREAS[(i - 1) % 5],
                gestor_id=(i - 1) % 5 + 1,
                data_inicio=date(2025, 1, 1),
                data_fim_prevista=date(2027, 12, 31),
            )
        )
    for i in range(1, 81):
        db.add(
            Fornecedor(
                id=i,
                razao_social=f"Fornecedor Fictício {i:02d} · {names[(i - 1) % 9]}",
                identificador_ficticio=f"FICTICIO-{i:04d}",
                categoria=["Serviços", "Materiais", "Tecnologia", "Logística"][(i - 1) % 4],
                email=f"fornecedor{i}@fornecedores.example",
            )
        )
    db.flush()
    for i in range(1, 1501):
        received = REFERENCE - timedelta(days=int(rng.integers(0, 270)))
        issued = received - timedelta(days=int(rng.integers(0, 12)))
        project = None if i % 11 == 0 else int(rng.integers(1, 46))
        responsible = (
            None if i % 17 == 0 else ((project or 1) - 1) % 5 + 11 + 5 * int(rng.integers(0, 11))
        )
        supplier = int(rng.integers(1, 81))
        number = f"DEMO-{i:06d}"
        if i % 50 == 0:
            source = db.get(Documento, i - 1)
            supplier, number, issued = (
                source.fornecedor_id,
                source.numero_documento,
                source.data_emissao,
            )
            received = max(received, issued)
        if i % 13 == 0:
            issued = received + timedelta(days=3)
        doc = Documento(
            id=i,
            numero_documento=number,
            fornecedor_id=supplier,
            projeto_id=project,
            responsavel_id=responsible,
            data_emissao=issued,
            data_recebimento=received,
            data_vencimento=max(received, issued) + timedelta(days=30),
            valor=Decimal(str(round(float(rng.lognormal(9.3, 1.1)), 2))),
            descricao="Recebimento sintético de serviços e materiais.",
            created_at=datetime.combine(received, datetime.min.time(), timezone.utc),
        )
        db.add(doc)
        db.flush()
        evaluate(db, doc, datetime.combine(received, datetime.min.time(), timezone.utc))
        audit(db, "documento", doc.id, "documento_recebido", novo=number)
    db.flush()
    issues = list(db.scalars(select(Pendencia)))
    for n, p in enumerate(issues):
        doc = db.get(Documento, p.documento_id)
        if n % 4 == 0 and p.data_abertura.date() <= REFERENCE - timedelta(days=2):
            if p.tipo == "classificacao":
                doc.projeto_id = 1
            elif p.tipo == "atribuicao":
                doc.responsavel_id = ((doc.projeto_id or 1) - 1) % 5 + 11
                p.responsavel_id = doc.responsavel_id
            elif p.tipo == "validacao":
                doc.data_emissao = doc.data_recebimento
            j = Justificativa(
                pendencia_id=p.id,
                usuario_id=p.responsavel_id or 7,
                texto="Justificativa sintética: informações conferidas e regularização registrada para demonstração.",
                status="aprovada",
                data_envio=p.data_abertura + timedelta(days=1),
                data_revisao=p.data_abertura + timedelta(days=2),
                comentario_revisor="Revisão demonstrativa aprovada.",
            )
            db.add(j)
            p.status = "encerrada"
            p.data_encerramento = p.data_abertura + timedelta(days=2)
            audit(db, "pendencia", p.id, "pendencia_encerrada", 2, "aberta", "encerrada")
            refresh_document(db, doc)
        elif n % 7 == 0:
            p.status = "em_tratamento"
            db.add(
                Justificativa(
                    pendencia_id=p.id,
                    usuario_id=p.responsavel_id or 7,
                    texto="Informações preliminares fictícias, aguardando documentação complementar.",
                    status="devolvida",
                    data_envio=p.data_abertura + timedelta(days=1),
                    comentario_revisor="Complemente a informação do projeto.",
                    data_revisao=p.data_abertura + timedelta(days=2),
                )
            )
        elif n % 9 == 0:
            p.status = "aguardando_revisao"
            db.add(
                Justificativa(
                    pendencia_id=p.id,
                    usuario_id=p.responsavel_id or 7,
                    texto="Análise demonstrativa concluída e enviada para revisão pelo responsável.",
                    data_envio=p.data_abertura + timedelta(days=1),
                )
            )
        if p.responsavel_id and n % 3 == 0:
            sent = p.data_abertura + timedelta(days=1)
            db.add(
                Cobranca(
                    pendencia_id=p.id,
                    destinatario_id=p.responsavel_id,
                    dia_envio=sent.date(),
                    data_envio=sent,
                    tipo="lembrete",
                    mensagem=f"Lembrete fictício: pendência #{p.id} aguarda tratamento.",
                )
            )
            audit(db, "pendencia", p.id, "cobranca_simulada", 1, novo="caixa_interna")
    db.flush()
    documents = {d.id: d for d in db.scalars(select(Documento))}
    pending = {p.id: p for p in db.scalars(select(Pendencia))}
    for d in documents.values():
        d.updated_at = d.created_at + timedelta(days=3)
    for event in db.scalars(select(Evento)):
        if event.entidade == "documento":
            event.timestamp = documents[event.entidade_id].created_at
        elif event.entidade == "pendencia":
            p = pending[event.entidade_id]
            event.timestamp = (
                p.data_encerramento
                if event.acao == "pendencia_encerrada"
                else p.data_abertura + timedelta(days=int(event.acao == "cobranca_simulada"))
            )
    db.flush()
    if db.get_bind().dialect.name == "postgresql":
        # Explicit seed IDs do not advance PostgreSQL SERIAL sequences automatically.
        for table in Base.metadata.sorted_tables:
            maximum = db.scalar(select(func.max(table.c.id)))
            if maximum:
                db.execute(
                    text(
                        "SELECT setval(pg_get_serial_sequence(:table_name, 'id'), :maximum, true)"
                    ),
                    {"table_name": table.name, "maximum": maximum},
                )
    db.commit()
    if export:
        export_dataset(db)


def export_dataset(db):
    folder = ROOT / "data" / "synthetic"
    folder.mkdir(parents=True, exist_ok=True)
    entities = {
        "projetos": Projeto,
        "fornecedores": Fornecedor,
        "usuarios": Usuario,
        "documentos_fiscais": Documento,
    }
    frames = {}
    for name, model in entities.items():
        rows = [
            {c.name: getattr(x, c.name) for c in model.__table__.columns}
            for x in db.scalars(select(model).order_by(model.id))
        ]
        frame = pd.DataFrame(rows)
        for column in frame.columns:
            if "data" in column or column.endswith("_at"):
                frame[column] = frame[column].map(
                    lambda x: x.isoformat() if hasattr(x, "isoformat") else x
                )
        for column in frame.columns:
            frame[column] = frame[column].map(
                lambda value: (
                    "'" + value
                    if isinstance(value, str)
                    and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
                    else value
                )
            )
        frame.to_csv(folder / f"{name}.csv", index=False, encoding="utf-8-sig")
        frames[name] = frame
    dictionary = [
        {
            "tabela": name,
            "campo": col.name,
            "tipo": str(col.type),
            "obrigatorio": not col.nullable,
            "observacao": "Dados exclusivamente sintéticos; identificadores sem validade fiscal.",
        }
        for name, model in entities.items()
        for col in model.__table__.columns
    ]
    frames["dicionario"] = pd.DataFrame(dictionary)
    with pd.ExcelWriter(folder / "dataset_demonstrativo.xlsx", engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="163C3D")
                cell.font = Font(color="FFFFFF", bold=True)
                ws.column_dimensions[cell.column_letter].width = 24
            table = Table(displayName=f"Tabela_{name}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    return Path(folder / "dataset_demonstrativo.xlsx")
